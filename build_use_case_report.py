from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"D:\YEAR3_HK1\KIENTAP_ - Copy\deliverables\Bao_cao_Use_Case_He_thong_Quan_ly_Thuong.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "16324F"
TEAL = "0F766E"
LIGHT_TEAL = "E8F5F3"
PALE = "F5F7FA"
BORDER = "D9E1EA"
TEXT = "25344A"
MUTED = "5E6F82"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def keep_with_next(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    pPr.append(el)

def set_run(run, size=None, bold=None, color=None, italic=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, size=10.5, bold=False, color=TEXT, italic=False):
    r = p.add_run(text)
    set_run(r, size, bold, color, italic)
    return r

def add_para(doc, text='', style=None, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.18
    if align is not None: p.alignment = align
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        add_text(p, item, 10.2)

def write_cell(cell, text, bold=False, color=TEXT, size=9.3, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None: p.alignment = align
    add_text(p, text, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    set_cell_border(cell)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        if widths: hdr.cells[i].width = Inches(widths[i])
        set_cell_shading(hdr.cells[i], NAVY)
        write_cell(hdr.cells[i], label, bold=True, color='FFFFFF', size=9.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if widths: cells[i].width = Inches(widths[i])
            if idx % 2 == 1: set_cell_shading(cells[i], PALE)
            write_cell(cells[i], str(val), size=9.1, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), TEAL); rPr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rPr.append(underline)
    rPr.append(OxmlElement('w:b'))
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def add_flow(doc, code, name, primary, precondition, steps, outcome, exceptions=None):
    h = doc.add_paragraph(style='Heading 2')
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    keep_with_next(h)
    add_text(h, f'{code}  {name}', size=13, bold=True, color='000000')
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    intro.paragraph_format.line_spacing = 1.1
    add_text(intro, 'Tác nhân chính: ', 9.7, True, MUTED)
    add_text(intro, primary, 9.7, True, TEAL)
    add_text(intro, '   |   Điều kiện: ', 9.7, True, MUTED)
    add_text(intro, precondition, 9.7, False, TEXT)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [('Mục tiêu', name), ('Luồng chính', None), ('Kết quả', outcome)]
    for label, value in labels:
        row = table.add_row().cells
        row[0].width = Inches(1.05); row[1].width = Inches(5.85)
        set_cell_shading(row[0], LIGHT_TEAL)
        write_cell(row[0], label, True, TEAL, 9.1)
        if label == 'Luồng chính':
            row[1].text = ''
            p = row[1].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for n, step in enumerate(steps, 1):
                if n > 1: p.add_run('\n')
                add_text(p, f'{n}. ', 9.1, True, TEAL)
                add_text(p, step, 9.1)
            set_cell_margin(row[1], 100, 120, 100, 120); set_cell_border(row[1])
        else:
            write_cell(row[1], value, False, TEXT, 9.1)
    if exceptions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
        add_text(p, 'Ngoại lệ/kiểm soát: ', 9.1, True, 'A14D2A')
        add_text(p, exceptions, 9.1, False, TEXT)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
for nm, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 13)]:
    st = styles[nm]; st.font.name = 'Aptos Display' if nm == 'Title' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
    st.font.size = Pt(size); st.font.color.rgb = RGBColor(0,0,0); st.font.bold = True
styles['Heading 1'].paragraph_format.space_before = Pt(14); styles['Heading 1'].paragraph_format.space_after = Pt(6)
title_pPr = styles['Title']._element.get_or_add_pPr()
for border in title_pPr.findall(qn('w:pBdr')):
    title_pPr.remove(border)

# Header and footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(header, 'HỆ THỐNG QUẢN LÝ VÀ ĐỀ XUẤT THƯỞNG', 8.5, True, MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(footer, 'Báo cáo use case  |  Bản mô tả phục vụ demo và bàn giao', 8.5, False, MUTED)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(78); p.paragraph_format.space_after = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(p, 'BÁO CÁO USE CASE', 10.5, True, TEAL)
p = doc.add_paragraph(style='Title')
p.paragraph_format.space_after = Pt(10)
add_text(p, 'Hệ thống Quản lý và Đề xuất Thưởng Thông minh', 24, True, '000000')
p = add_para(doc, 'Các luồng nghiệp vụ từ tính thưởng, đề xuất, phê duyệt đến chi trả', after=22)
for r in p.runs: set_run(r, 13, False, MUTED)

meta = doc.add_table(rows=4, cols=2); meta.alignment = WD_TABLE_ALIGNMENT.LEFT; meta.autofit = False
for i, (label, val) in enumerate([
    ('Mục đích', 'Giúp người xem hiểu nhanh cách hệ thống vận hành theo vai trò và trách nhiệm.'),
    ('Phạm vi', 'Đăng nhập, tổng quan, đề xuất thưởng, phê duyệt, tối ưu cơ chế, kiểm tra dữ liệu và chi trả.'),
    ('Đối tượng đọc', 'Nhóm dự án, giảng viên hoặc đơn vị đánh giá, quản lý nghiệp vụ C&B.'),
    ('Phiên bản', 'Bản mô phỏng web hiện hành - tháng 09 năm 2026.'),
]):
    meta.rows[i].cells[0].width=Inches(1.25); meta.rows[i].cells[1].width=Inches(5.65)
    set_cell_shading(meta.rows[i].cells[0], LIGHT_TEAL)
    write_cell(meta.rows[i].cells[0], label, True, TEAL, 9.5)
    write_cell(meta.rows[i].cells[1], val, False, TEXT, 9.5)

doc.add_page_break()

# Overview
h=doc.add_paragraph(style='Heading 1'); add_text(h, '1. Mục tiêu và phạm vi vận hành', 16, True, '000000')
add_para(doc, 'Hệ thống quản lý toàn bộ vòng đời thưởng theo kỳ. Thay vì chỉ trả về một con số, hệ thống kết hợp dữ liệu hiệu suất, quy chế, dự báo lịch sử và ý kiến quản lý để tạo đề xuất có thể giải thích. Quyết định chi trả chỉ được xác lập sau khi kiểm soát dữ liệu và phê duyệt đúng thẩm quyền.', after=7)

h=doc.add_paragraph(style='Heading 2'); add_text(h, 'Vòng đời nghiệp vụ', 13, True, '000000')
add_table(doc, ['Giai đoạn', 'Mục đích', 'Kết quả bàn giao'], [
    ('1. Ghi nhận', 'Chọn kỳ, đồng bộ hồ sơ, đo hiệu suất theo trọng số.', 'Dữ liệu đầu vào và chỉ số tổng quan.'),
    ('2. Đề xuất', 'Hệ thống tính mức thưởng; quản lý xem xét hoặc điều chỉnh có lý do.', 'Hồ sơ chờ C&B phê duyệt.'),
    ('3. Kiểm soát', 'C&B kiểm tra đề xuất, cảnh báo và tác động ngân sách.', 'Quyết định phê duyệt hoặc từ chối.'),
    ('4. Chi trả', 'Tổng hợp hồ sơ đã duyệt và xuất payroll.', 'Danh sách chi trả theo kỳ.'),
], widths=[1.25, 3.2, 2.45])

h=doc.add_paragraph(style='Heading 2'); add_text(h, 'Ma trận quyền', 13, True, '000000')
add_table(doc, ['Vai trò', 'Phạm vi dữ liệu', 'Thao tác chính'], [
    ('Admin C&B và Tài chính', 'Toàn công ty.', 'Theo dõi kỳ, duyệt/từ chối đề xuất, xử lý cảnh báo, duyệt cơ chế và xuất payroll.'),
    ('Quản lý trực tiếp', 'Nhân sự thuộc phòng ban phụ trách.', 'Xem hiệu suất, chấp nhận/điều chỉnh mức thưởng, gửi đề xuất và đề xuất tối ưu cơ chế.'),
    ('Nhân viên', 'Duy nhất hồ sơ cá nhân.', 'Xem phiếu thưởng, các thành phần tính và gửi yêu cầu giải trình.'),
], widths=[1.5, 2.0, 3.4])

add_para(doc, 'Nguyên tắc kiểm soát: nhân viên không truy cập được màn hình quản trị; quản lý không xuất payroll; chỉ Admin C&B có quyền phê duyệt cấp hai, thay đổi trạng thái cuối cùng và xuất chi trả.', after=5)

doc.add_page_break()

# Flows 1-3
h=doc.add_paragraph(style='Heading 1'); add_text(h, '2. Use case theo flow', 16, True, '000000')
add_flow(doc, 'UC 01', 'Đăng nhập và giới hạn quyền truy cập', 'Admin C&B, Quản lý, Nhân viên', 'Người dùng có tài khoản mẫu hợp lệ.', [
    'Chọn một trong ba vai trò mẫu hoặc nhập tên đăng nhập và mật khẩu.',
    'Hệ thống xác thực, hiển thị tên vai trò và dựng menu theo quyền.',
    'Người dùng mở màn hình được phép; dữ liệu được lọc theo phạm vi quyền.'
], 'Mỗi vai trò chỉ nhìn thấy dữ liệu và chức năng đúng thẩm quyền.', 'Nếu nhập sai thông tin, hệ thống báo lỗi. Nếu nhân viên cố truy cập route quản trị, hệ thống chuyển về khu vực cá nhân.')

add_flow(doc, 'UC 02', 'Theo dõi kỳ thưởng và diễn giải số liệu', 'Admin C&B hoặc Quản lý', 'Đã đăng nhập với quyền quản trị/điều hành.', [
    'Chọn kỳ tính và mốc ngày công cần theo dõi.',
    'Hệ thống cập nhật chỉ số theo dữ liệu của kỳ: tỷ lệ đạt có trọng số, thưởng tạm tính, sử dụng quỹ và tiến độ thời gian.',
    'Người dùng bấm “Cách tính” để xem công thức cùng số liệu thực tế đang áp dụng.'
], 'Quản lý có một bức tranh điều hành thống nhất theo kỳ và có thể đối chiếu nguồn số.', 'Bộ lọc phòng ban làm thay đổi tập dữ liệu được tổng hợp; số tiền có thể được ẩn khi thuyết trình.')

add_flow(doc, 'UC 03', 'Quản lý đề xuất hoặc điều chỉnh mức thưởng', 'Quản lý trực tiếp', 'Đã chọn hồ sơ nhân sự thuộc đơn vị quản lý.', [
    'Mở hồ sơ nhân sự để xem hiệu suất, ba nguồn cấu thành và công thức đề xuất.',
    'Chọn đồng ý mức hệ thống hoặc nhập mức điều chỉnh trên thanh What If, kèm lý do khi cần.',
    'Gửi đề xuất; hệ thống lưu mức, số tiền tương ứng, người gửi và thời điểm gửi.'
], 'Hồ sơ chuyển sang trạng thái chờ C&B phê duyệt.', 'Quản lý chỉ thao tác trên nhân sự thuộc quyền. Sai lệch đáng kể so với mức hệ thống phải có căn cứ và chuyển duyệt cấp hai.')

doc.add_page_break()

# Flows 4-6
add_flow(doc, 'UC 04', 'C&B phê duyệt hoặc từ chối đề xuất thưởng', 'Admin C&B và Tài chính', 'Có hồ sơ ở trạng thái chờ phê duyệt.', [
    'Mở hàng đợi duyệt để xem mức hệ thống, mức quản lý đề xuất, lý do và tác động tài chính.',
    'Chọn phê duyệt để chốt mức thưởng; hoặc chọn từ chối và nhập lý do phản hồi.',
    'Hệ thống cập nhật trạng thái, người duyệt, ngày duyệt hoặc lý do từ chối.'
], 'Hồ sơ được duyệt sẵn sàng cho chi trả, hoặc quay lại vòng đề xuất với phản hồi rõ ràng.', 'Chỉ Admin C&B được thay đổi quyết định cuối cùng; lý do từ chối được lưu để quản lý bổ sung.')

add_flow(doc, 'UC 05', 'Đề xuất tối ưu cơ chế thưởng', 'Quản lý khu vực', 'Quản lý đã đăng nhập và nhận diện được vấn đề cơ chế tại đơn vị.', [
    'Mở chức năng tạo đề xuất cơ chế, nhập tiêu đề, nhóm vấn đề, mô tả và tác động dự kiến.',
    'Gửi đề xuất đến C&B; thẻ đề xuất xuất hiện trên khu vực tổng quan của Admin.',
    'Admin xem xét từng thẻ hoặc duyệt hàng loạt khi phù hợp.'
], 'Đề xuất cơ chế có trạng thái theo dõi và luồng xét duyệt tập trung.', 'Quản lý chỉ thấy đề xuất của đơn vị mình; Admin là bên duy nhất phê duyệt hoặc từ chối.')

add_flow(doc, 'UC 06', 'Kiểm tra và xử lý bất thường dữ liệu', 'Admin C&B và Tài chính', 'Hệ thống phát hiện hồ sơ cần kiểm tra.', [
    'Mở hàng đợi cảnh báo dữ liệu để xem nguyên nhân: trùng hợp đồng, thiếu chỉ tiêu, chưa gán cơ chế, tỷ lệ bất thường hoặc dưới sàn.',
    'Chọn hành động khắc phục tương ứng, ví dụ khấu trừ bản ghi trùng, gán chỉ tiêu/cơ chế, áp trần hoặc áp sàn.',
    'Hệ thống tính lại dữ liệu liên quan và làm mới hàng đợi.'
], 'Hồ sơ được đưa về dữ liệu hợp lệ trước khi đi vào chi trả.', 'Các hành động khắc phục chỉ dành cho Admin để tránh tự ý thay đổi dữ liệu nguồn.')

doc.add_page_break()

# Flows 7/8 and calculation notes
add_flow(doc, 'UC 07', 'Nhân viên xem phiếu thưởng và yêu cầu giải trình', 'Nhân viên', 'Đăng nhập bằng tài khoản nhân viên.', [
    'Mở “Thưởng của tôi” để xem số tiền, tỷ lệ hoàn thành và trạng thái hồ sơ.',
    'Đọc phần diễn giải: chỉ tiêu, thực đạt, yếu tố hiệu chỉnh và mức thưởng.',
    'Gửi yêu cầu giải trình nếu cần làm rõ quyết định.'
], 'Nhân viên có thông tin minh bạch về chính hồ sơ của mình mà không truy cập dữ liệu người khác.', 'Thông tin hiển thị theo quyền cá nhân; yêu cầu giải trình tạo đầu vào để bộ phận phụ trách tiếp nhận.')

add_flow(doc, 'UC 08', 'Tổng hợp và xuất payroll', 'Admin C&B và Tài chính', 'Hồ sơ đã được phê duyệt trong kỳ.', [
    'Mở màn hình tính thưởng và chi trả.',
    'Hệ thống tổng hợp hồ sơ đã duyệt, tổng tiền và số hồ sơ còn chờ.',
    'Admin xuất tệp Payroll CSV cho bước xử lý tài chính tiếp theo.'
], 'Có danh sách chi trả theo kỳ, chỉ bao gồm hồ sơ đủ điều kiện.', 'Quản lý và nhân viên không có quyền xuất dữ liệu chi trả tài chính.')

h=doc.add_paragraph(style='Heading 1'); add_text(h, '3. Quy tắc nghiệp vụ cần nhớ', 16, True, '000000')
add_table(doc, ['Quy tắc', 'Diễn giải sử dụng trong hệ thống'], [
    ('Ba nguồn đề xuất', 'Mức hệ thống tổng hợp 70% quy chế, 30% dự báo lịch sử và phần điều chỉnh đánh giá của quản lý.'),
    ('Mức hiển thị', 'Danh sách nhân sự và màn hình chi tiết cùng hiển thị mức thưởng đang áp dụng; mức hệ thống được giữ riêng để đối chiếu.'),
    ('Công thức có dữ liệu thật', 'Các cửa sổ “Cách tính” lấy số theo kỳ, ngày công, phòng ban và quyền người dùng đang chọn; không dùng số viết cứng trong giao diện.'),
    ('Phê duyệt', 'Đề xuất từ quản lý không trở thành mức chi trả cuối cùng cho đến khi Admin C&B phê duyệt.'),
    ('Tính nhất quán', 'Khi thay đổi kỳ hoặc bộ lọc, tổng quan, danh sách và các chỉ số liên quan được tính lại từ cùng nguồn dữ liệu.'),
], widths=[1.65, 5.25])

h=doc.add_paragraph(style='Heading 1'); add_text(h, '4. Kịch bản demo ngắn', 16, True, '000000')
add_bullets(doc, [
    'Đăng nhập Admin, chọn kỳ và bấm “Cách tính” để chứng minh số liệu có nguồn.',
    'Chuyển sang Quản lý, chọn một nhân sự, điều chỉnh mức thưởng và gửi C&B duyệt.',
    'Quay lại Admin, mở hàng đợi, phê duyệt hoặc từ chối có lý do.',
    'Chuyển sang Nhân viên để minh họa giới hạn quyền và phiếu thưởng minh bạch.',
    'Trở lại Admin, xử lý một cảnh báo dữ liệu và xuất payroll sau khi hồ sơ được duyệt.'
])

doc.add_page_break()

# Client-oriented expanded guide
h=doc.add_paragraph(style='Heading 1'); add_text(h, '5. Hướng dẫn đọc và đối chiếu trên website', 16, True, '000000')
add_para(doc, 'Phần này viết cho người yêu cầu dự án. Mỗi flow trả lời năm câu hỏi: dùng khi nào, ai thao tác, thao tác ở đâu, hệ thống phản hồi gì và người quản lý nhận được giá trị gì. Có thể mở website trong lúc đọc để đối chiếu từng màn hình.', after=6)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5)
add_text(p, 'Website đang chạy: ', 10.4, True, TEXT); add_hyperlink(p, 'incentive-management-demo.vercel.app', 'https://incentive-management-demo.vercel.app/')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(10)
add_text(p, 'Mã nguồn dự án: ', 10.4, True, TEXT); add_hyperlink(p, 'github.com/hongtrachy/kien_tap_2026', 'https://github.com/hongtrachy/kien_tap_2026')

h=doc.add_paragraph(style='Heading 2'); add_text(h, 'Cách trải nghiệm nhanh', 13, True, '000000')
add_table(doc, ['Bước', 'Người xem cần làm', 'Điều cần quan sát'], [
    ('1', 'Mở website và chọn thẻ vai trò mẫu ở trang đăng nhập.', 'Tên người dùng, menu bên trái và dữ liệu được thay đổi theo vai trò.'),
    ('2', 'Bắt đầu bằng Admin để xem tổng quan kỳ.', 'Bộ lọc kỳ, các thẻ chỉ số và các đề xuất cơ chế đang chờ duyệt.'),
    ('3', 'Chuyển sang Quản lý, mở Đề xuất mức thưởng.', 'Danh sách nhân sự thuộc quyền, ba nguồn dữ liệu và nút gửi duyệt.'),
    ('4', 'Chuyển sang Nhân viên, mở Thưởng của tôi.', 'Chỉ nhìn thấy hồ sơ cá nhân và phần giải trình dễ hiểu.'),
], widths=[0.55, 3.1, 3.25])

doc.add_page_break()
h=doc.add_paragraph(style='Heading 1'); add_text(h, '6. Diễn giải chi tiết từng flow', 16, True, '000000')

def add_client_flow(title, why, where, watch, example, business_value, evidence):
    h=doc.add_paragraph(style='Heading 2'); h.paragraph_format.space_before=Pt(10); keep_with_next(h)
    add_text(h, title, 13, True, '000000')
    for lead, value in [('Vì sao cần flow này', why), ('Thao tác trên web', where), ('Hệ thống hiển thị và kiểm soát', watch), ('Ví dụ minh họa', example), ('Giá trị cho bên yêu cầu', business_value), ('Minh chứng có thể đối chiếu', evidence)]:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.12
        add_text(p, lead + ': ', 9.8, True, TEAL)
        add_text(p, value, 9.8, False, TEXT)

add_client_flow(
    'Flow A  Đăng nhập và xem đúng phạm vi quyền',
    'Một hệ thống thưởng có dữ liệu nhạy cảm. Người xem cần thấy ngay rằng nhân viên không thể xem bảng lương chung, còn quản lý chỉ xem người thuộc bộ phận mình.',
    'Tại trang đầu, bấm một trong ba thẻ Admin, Quản lý hoặc Nhân viên. Hệ thống có sẵn tài khoản demo để không cần tạo dữ liệu mới khi trình bày.',
    'Thanh điều hướng thay đổi theo vai trò. Nếu nhân viên cố mở màn hình quản trị, hệ thống chặn và đưa về khu vực cá nhân. Admin xem toàn công ty; Quản lý xem đơn vị mình; Nhân viên chỉ xem một hồ sơ.',
    'Chọn Nhân viên rồi quan sát: không có mục xuất Payroll, không có hàng đợi phê duyệt và không có danh sách nhân sự của người khác.',
    'Giảm rủi ro lộ dữ liệu và làm rõ trách nhiệm xử lý. Đây là điều kiện nền để các flow tiếp theo có ý nghĩa.',
    'Có thể kiểm chứng bằng cách chuyển vai trò trên thanh trên cùng hoặc đăng xuất rồi chọn lại vai trò mẫu.'
)

add_client_flow(
    'Flow B  Theo dõi một kỳ thưởng bằng số liệu có thể giải thích',
    'Người quản lý cần biết kỳ thưởng đang đi tới đâu, quỹ còn an toàn hay không và vì sao con số thay đổi; không chỉ cần một dashboard đẹp.',
    'Đăng nhập Admin hoặc Quản lý, vào Tổng quan kỳ thưởng. Chọn kỳ tính, phòng ban nếu cần và mốc ngày công.',
    'Bốn thẻ chính cập nhật đồng bộ: tỷ lệ đạt có trọng số, thưởng tạm tính, tỷ lệ sử dụng quỹ và tiến độ thời gian. Nút Cách tính mở cửa sổ diễn giải với công thức và giá trị của đúng kỳ/mốc đang chọn.',
    'Nếu chọn ngày công 15/22, hệ thống cho biết tiến độ thời gian là 68,2%. Khi đổi kỳ hoặc phòng ban, tổng quan và danh sách liên quan cùng tính lại từ tập dữ liệu đang lọc.',
    'Người yêu cầu dự án có thể kiểm tra tính minh bạch: số liệu không được đặt cố định trên HTML/CSS mà lấy từ phép tính nghiệp vụ.',
    'Bấm Cách tính ở từng thẻ. Ví dụ thẻ Thưởng tạm tính hiển thị tổng thưởng nhân với tỷ lệ tiến độ thực tế thay vì hiển thị 0 đồng hoặc một số ngẫu nhiên.'
)

add_client_flow(
    'Flow C  Hệ thống tạo đề xuất thưởng thông minh cho từng nhân sự',
    'Quyết định thưởng không nên chỉ dựa vào cảm tính hoặc một chỉ số KPI. Flow này giúp người quản lý nhìn được cả quy chế, lịch sử và đánh giá thực tế.',
    'Đăng nhập Quản lý, vào Đề xuất mức thưởng, rồi chọn một nhân sự trong danh sách bên trái.',
    'Màn hình chi tiết tách rõ ba nguồn: mức theo quy chế, dự báo lịch sử và điều chỉnh của quản lý. Bên dưới là công thức hai bước, số tiền quy đổi và giải thích bằng tiếng Việt.',
    'Với Lê Hoàng Chi, danh sách hiển thị 105,0% mức thưởng đang chờ duyệt. Trong chi tiết, mức quản lý đề xuất là 105,0%, còn mức hệ thống đề xuất theo công thức là 108,0%. Hai giá trị không bị đánh đồng.',
    'Người xem hiểu được mức nào là khuyến nghị của hệ thống và mức nào là quyết định nghiệp vụ của quản lý. Điều này đặc biệt quan trọng khi có chênh lệch cần được giải trình.',
    'Chọn Lê Hoàng Chi và so sánh khu vực đầu trang với ô công thức. Phần giải thích nêu rõ mức quản lý trình duyệt, mức hệ thống và lý do điều chỉnh.'
)

add_client_flow(
    'Flow D  Quản lý điều chỉnh và gửi C B duyệt',
    'Máy tính có thể đưa ra đề xuất tốt nhưng không thay thế được ngữ cảnh địa bàn, khách hàng hoặc sự kiện thực tế. Quản lý cần được quyền đưa ra ý kiến có kiểm soát.',
    'Trong hồ sơ nhân sự, quản lý chọn Đồng ý mức đề xuất hoặc dùng thanh What If để nhập một mức khác, bổ sung lý do và bấm Xác nhận gửi duyệt.',
    'Hệ thống cập nhật mức thưởng, số tiền tương ứng, người gửi, ngày gửi và trạng thái Chờ duyệt. Nếu mức điều chỉnh lệch đáng kể, giao diện yêu cầu căn cứ thay vì cho gửi một quyết định không giải thích được.',
    'Quản lý đề xuất 105% trong khi hệ thống tính 108%. Mức 105% xuất hiện nhất quán ở danh sách và detail, nhưng công thức 108% vẫn được giữ để C&B đánh giá chênh lệch.',
    'Giữ quyền đánh giá nghiệp vụ ở tuyến quản lý nhưng tạo dấu vết đầy đủ cho bước phê duyệt sau đó.',
    'Sau khi gửi, vào hàng đợi phê duyệt bằng tài khoản Admin để thấy cùng hồ sơ, mức đề xuất và căn cứ đi kèm.'
)

add_client_flow(
    'Flow E  C B phê duyệt hoặc từ chối với lý do rõ ràng',
    'Chi trả thưởng cần một điểm kiểm soát độc lập về ngân sách và tính công bằng. Admin C&B là vai trò chịu trách nhiệm chốt quyết định cuối.',
    'Đăng nhập Admin, vào Kiểm tra dữ liệu hoặc hàng đợi đề xuất. Mở hồ sơ đang chờ, đọc mức hệ thống, mức quản lý trình duyệt, lý do và tác động tiền thưởng.',
    'Admin bấm Phê duyệt để chốt hồ sơ hoặc Từ chối để nhập phản hồi. Hệ thống lưu người duyệt, ngày duyệt hoặc lý do từ chối; sau đó làm mới các màn hình liên quan.',
    'Nếu thiếu căn cứ, Admin có thể ghi “Cần đính kèm số liệu so sánh cùng kỳ và căn cứ điều chỉnh”. Quản lý biết chính xác cần bổ sung gì thay vì chỉ nhận trạng thái bị từ chối.',
    'Giảm quyết định cảm tính và giúp bên yêu cầu chứng minh rằng mọi ngoại lệ đều có người chịu trách nhiệm phê duyệt.',
    'Quan sát trạng thái hồ sơ chuyển từ Chờ duyệt sang Đã duyệt hoặc Bị từ chối; phiếu thưởng và payroll chỉ dùng hồ sơ được duyệt.'
)

add_client_flow(
    'Flow F  Quản lý đề xuất thay đổi cơ chế thưởng',
    'Ngoài điều chỉnh từng người, doanh nghiệp đôi khi cần thay đổi cơ chế cho cả khu vực, ví dụ do thị trường biến động hoặc có nhóm nhân sự tập trung sát một mốc thưởng.',
    'Quản lý mở khu vực Đề xuất tối ưu cơ chế, nhập nội dung đề xuất và gửi C&B xét duyệt.',
    'Đề xuất được tạo thành thẻ có tiêu đề, mô tả, khu vực gửi, người gửi, ngày gửi và tác động dự kiến. Thẻ này tự hiện trên dashboard của Admin để không bị thất lạc qua trao đổi ngoài hệ thống.',
    'Một quản lý có thể đề xuất hiệu chỉnh hệ số độ khó cho vùng đang suy giảm sức mua. Admin nhìn thấy cả tác động tăng/giảm quỹ trước khi duyệt.',
    'Biến kiến nghị nghiệp vụ thành dữ liệu có thể theo dõi, thay vì chỉ tồn tại trong email hoặc trao đổi miệng.',
    'Đăng nhập Quản lý để gửi thử; sau đó chuyển Admin và xem thẻ mới trong Tổng quan kỳ thưởng.'
)

add_client_flow(
    'Flow G  Phát hiện và xử lý dữ liệu bất thường trước khi trả thưởng',
    'Một con số thưởng đúng công thức vẫn sai nếu dữ liệu đầu vào có hợp đồng trùng, chỉ tiêu thiếu hoặc tỷ lệ không hợp lý.',
    'Admin vào Kiểm tra dữ liệu để xem hàng đợi cảnh báo. Từ từng cảnh báo, chọn thao tác khắc phục phù hợp.',
    'Hệ thống chỉ ra loại lỗi và hành động cụ thể: khấu trừ bản ghi trùng, gán chỉ tiêu, gán cơ chế, áp trần hoặc áp sàn. Sau khi xử lý, số liệu liên quan được tính lại.',
    'Phạm Tiến Dũng có hợp đồng trùng 90 triệu đồng. Khi Admin khấu trừ bản ghi này, doanh số được đưa về đúng mức; tỷ lệ và thưởng được cập nhật theo dữ liệu đã làm sạch.',
    'Ngăn việc chi trả dựa trên dữ liệu lỗi, đồng thời làm cho quá trình kiểm tra có thể kiểm toán lại.',
    'Mở cảnh báo của Phạm Tiến Dũng để xem nhãn hợp đồng trùng và nút khấu trừ bản ghi trùng.'
)

add_client_flow(
    'Flow H  Nhân viên tự xem và yêu cầu giải trình',
    'Minh bạch không chỉ phục vụ bộ phận quản trị. Người nhận thưởng cần hiểu mức của mình đến từ đâu và có kênh yêu cầu làm rõ.',
    'Đăng nhập Nhân viên, vào Thưởng của tôi hoặc Phiếu thưởng cá nhân.',
    'Hệ thống hiển thị hồ sơ cá nhân, tỷ lệ đạt, mức thưởng và phần diễn giải theo ngôn ngữ dễ đọc. Nhân viên có thể mở yêu cầu giải trình mà không nhìn thấy dữ liệu của đồng nghiệp.',
    'Trần Thị Bình thấy chỉ tiêu, thực đạt, hệ số điều chỉnh do thị trường và mức thưởng cuối. Phần này giải thích quan hệ giữa dữ liệu công việc và tiền thưởng thay vì chỉ hiển thị một dòng tiền.',
    'Giảm thắc mắc lặp lại, tăng niềm tin và tạo kênh phản hồi có cấu trúc.',
    'Thử chuyển qua tài khoản Nhân viên và quan sát menu rút gọn cùng nút Gửi yêu cầu giải trình.'
)

add_client_flow(
    'Flow I  Chỉ xuất payroll sau khi hồ sơ đủ điều kiện',
    'Lô chi trả cần loại trừ hồ sơ chưa duyệt hoặc có vấn đề dữ liệu, nếu không rủi ro tài chính sẽ chuyển sang bộ phận kế toán.',
    'Admin mở Tính thưởng và chi trả, kiểm tra tổng số hồ sơ đã duyệt và số hồ sơ còn chờ, rồi chọn Xuất file Payroll CSV.',
    'Tổng tiền, số hồ sơ và trạng thái lô được hiển thị trước khi xuất. Chỉ Admin có nút này; Quản lý và Nhân viên không truy cập được.',
    'Nếu còn hồ sơ chờ C&B xét duyệt, phần tổng hợp cho biết rõ số lượng; người phụ trách không nhầm đây là lô đã chốt hoàn toàn.',
    'Tạo một điểm giao nhận rõ ràng giữa nghiệp vụ C&B và bước chi trả tài chính.',
    'Đăng nhập Admin, mở Tính thưởng và chi trả, đối chiếu tổng tiền với số hồ sơ được phê duyệt trước khi tải CSV.'
)

h=doc.add_paragraph(style='Heading 1'); add_text(h, '7. Minh chứng nghiệp vụ và giới hạn hiện tại', 16, True, '000000')
add_para(doc, 'Báo cáo này mô tả bản demo vận hành bằng dữ liệu mô phỏng có kiểm soát. Các phép tính, trạng thái và phân quyền đều chạy trực tiếp trong website để phục vụ đánh giá flow. Cơ chế đăng nhập hiện là tài khoản mẫu trên trình duyệt; khi đưa vào môi trường doanh nghiệp, cần thay bằng xác thực máy chủ, phân quyền theo danh tính thực và lưu vết giao dịch tập trung.', after=6)
add_table(doc, ['Nội dung cần chứng minh', 'Cách kiểm tra trên website', 'Ý nghĩa cho bên yêu cầu'], [
    ('Tính minh bạch số liệu', 'Dùng nút Cách tính và xem công thức theo kỳ/ngày công đang chọn.', 'Xác nhận số liệu có nguồn và có thể giải thích.'),
    ('Nhất quán mức thưởng', 'Chọn Lê Hoàng Chi và so sánh list, detail, mức hệ thống và mức quản lý.', 'Phân biệt rõ KPI, khuyến nghị hệ thống và mức đang trình duyệt.'),
    ('Tách bạch thẩm quyền', 'Đổi vai trò từ Admin sang Quản lý rồi Nhân viên.', 'Xác nhận đúng người đúng việc, không lộ dữ liệu.'),
    ('Kiểm soát ngoại lệ', 'Mở hàng đợi cảnh báo và xử lý ca hợp đồng trùng.', 'Xác nhận dữ liệu sai không đi thẳng sang chi trả.'),
    ('Sẵn sàng chi trả', 'Vào Payroll sau khi duyệt hồ sơ.', 'Xác nhận chỉ xuất lô từ hồ sơ đạt điều kiện.'),
], widths=[1.55, 3.15, 2.2])

h=doc.add_paragraph(style='Heading 1'); add_text(h, '8. Kết luận cho người yêu cầu dự án', 16, True, '000000')
add_para(doc, 'Website không chỉ tính một con số thưởng. Nó đưa người dùng đi qua một chuỗi quyết định có kiểm soát: nhìn dữ liệu, tạo đề xuất, cho phép quản lý giải thích ngoại lệ, để C&B phê duyệt độc lập, kiểm tra dữ liệu trước chi trả và trả lại diễn giải minh bạch cho nhân viên. Khi demo, nên đi lần lượt theo Flow B, C, D, E, H và I để người xem thấy toàn bộ vòng đời trong một mạch liên tục.', after=4)

doc.save(OUT)
print(OUT)
